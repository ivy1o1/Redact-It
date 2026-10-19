"""Tests for PDF and DOCX text extraction.

The contract under test: every extracted span can be re-located exactly (word
boxes for PDF, run spans for DOCX), and every document this release cannot fully
handle - scans, encrypted files, image-only pages, uncovered Word constructs -
says so instead of looking like a success.
"""

from __future__ import annotations

import io
from pathlib import Path

import docx
import pymupdf
import pytest
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from app.errors import (
    EmptyDocumentError,
    EncryptedDocumentError,
    ExtractionBudgetError,
    ExtractionError,
    ScannedDocumentError,
    UnsupportedFileTypeError,
)
from app.extraction import DOCX_LIMITATIONS, PDF_LIMITATIONS, extract_document

from .conftest import (
    REPORT_DOCX_PARAGRAPHS,
    REPORT_LINES,
    TABLE_ROW,
    build_docx,
    build_docx_with_split_runs,
    build_encrypted_pdf,
    build_pdf,
    build_scanned_pdf,
)

MAX_CHARS = 100_000


def extract(path: Path, *, name: str | None = None, max_chars: int = MAX_CHARS):
    kind = "pdf" if path.suffix.lower() == ".pdf" else "docx"
    return extract_document(
        path,
        file_kind=kind,
        document_id="doc_test",
        document_name=name or path.name,
        max_chars=max_chars,
    )


@pytest.fixture
def pdf_path(tmp_path: Path) -> Path:
    path = tmp_path / "report.pdf"
    path.write_bytes(build_pdf())
    return path


@pytest.fixture
def docx_path(tmp_path: Path) -> Path:
    path = tmp_path / "report.docx"
    path.write_bytes(build_docx(table_row=TABLE_ROW, header_text="Confidential: Alice Johnson"))
    return path


class TestPdfExtraction:
    def test_metadata_describes_the_file(self, pdf_path: Path) -> None:
        document = extract(pdf_path, name="report.pdf")
        assert document.file_kind == "pdf"
        assert document.document_name == "report.pdf"
        assert document.document_id == "doc_test"
        assert document.page_count == 1
        assert document.char_count == sum(len(s.text) for s in document.segments)
        assert document.has_text_layer is True
        assert document.scanned is False
        assert document.encrypted is False
        assert document.truncated is False
        assert document.size_bytes == pdf_path.stat().st_size

    def test_sha256_covers_the_whole_file(self, pdf_path: Path) -> None:
        import hashlib

        document = extract(pdf_path)
        assert document.sha256 == hashlib.sha256(pdf_path.read_bytes()).hexdigest()

    def test_one_segment_per_page_with_one_based_pages(self, tmp_path: Path) -> None:
        path = tmp_path / "three.pdf"
        path.write_bytes(build_pdf(pages=3))
        document = extract(path)
        assert document.page_count == 3
        assert [s.page for s in document.segments] == [1, 2, 3]
        assert [s.segment_index for s in document.segments] == [0, 1, 2]

    def test_extracted_text_contains_the_synthetic_content(self, pdf_path: Path) -> None:
        document = extract(pdf_path)
        text = document.segments[0].text
        for fragment in ("Alice Johnson", "(555) 123-4567", "42 Maple Street", "Algebra II"):
            assert fragment in text

    def test_word_offsets_slice_back_to_the_word_text(self, pdf_path: Path) -> None:
        segment = extract(pdf_path).segments[0]
        assert segment.words
        for word in segment.words:
            assert segment.text[word.char_start : word.char_end] == word.text

    def test_words_cover_the_text_without_gaps_or_overlaps(self, pdf_path: Path) -> None:
        segment = extract(pdf_path).segments[0]
        assert segment.words[0].char_start == 0
        for previous, current in zip(segment.words, segment.words[1:]):
            # Words are joined by a single space in the extracted text.
            assert current.char_start == previous.char_end + 1

    def test_word_boxes_are_real_rectangles_on_the_right_page(self, pdf_path: Path) -> None:
        segment = extract(pdf_path).segments[0]
        for word in segment.words:
            x0, y0, x1, y1 = word.rect
            assert x1 > x0 and y1 > y0
            assert word.page == 1

    def test_words_in_range_returns_exactly_the_words_of_a_match(self, pdf_path: Path) -> None:
        document = extract(pdf_path)
        segment = document.segments[0]
        start = segment.text.index("Alice Johnson")
        boxes = document.words_in_range(0, start, start + len("Alice Johnson"))
        assert [b.text for b in boxes] == ["Alice", "Johnson"]

    def test_words_in_range_is_empty_for_a_range_that_covers_nothing(self, pdf_path: Path) -> None:
        document = extract(pdf_path)
        assert document.words_in_range(0, 0, 0) == []
        assert document.words_in_range(99, 0, 50) == []

    def test_line_ranges_cover_the_page(self, tmp_path: Path) -> None:
        path = tmp_path / "lines.pdf"
        path.write_bytes(build_pdf(lines=["First line here", "Second line there"]))
        segment = extract(path).segments[0]
        assert len(segment.line_ranges) == 2
        assert segment.line_ranges[0][0] == 0
        assert segment.line_ranges[-1][1] == len(segment.text)
        for start, end in segment.line_ranges:
            assert 0 <= start < end <= len(segment.text)

    def test_a_page_shared_between_lines_is_still_one_segment(self, pdf_path: Path) -> None:
        segment = extract(pdf_path).segments[0]
        assert len(segment.line_ranges) == len(REPORT_LINES)

    def test_pdf_segments_have_no_paragraph_path(self, pdf_path: Path) -> None:
        assert all(s.paragraph_path is None for s in extract(pdf_path).segments)

    def test_limitations_are_declared(self, pdf_path: Path) -> None:
        document = extract(pdf_path)
        assert document.limitations == PDF_LIMITATIONS
        assert any("No OCR" in line for line in document.limitations)
        assert any("word granularity" in line for line in document.limitations)

    def test_a_text_page_next_to_an_image_only_page_warns_about_the_image_page(
        self, tmp_path: Path
    ) -> None:
        document = pymupdf.open()
        text_page = document.new_page()
        text_page.insert_textbox(pymupdf.Rect(50, 50, 560, 780), "Text page", fontsize=11)
        image_page = document.new_page()
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 100, 80), 1)
        pixmap.clear_with(180)
        image_page.insert_image(pymupdf.Rect(50, 50, 150, 130), pixmap=pixmap)
        path = tmp_path / "mixed.pdf"
        path.write_bytes(document.tobytes())
        document.close()

        extracted = extract(path)
        assert extracted.scanned is False
        assert [s.page for s in extracted.segments] == [1]
        assert any("no text layer" in w and "[2]" in w for w in extracted.warnings)


class TestPdfFailures:
    def test_an_image_only_pdf_is_refused_as_a_scan(self, tmp_path: Path) -> None:
        path = tmp_path / "scan.pdf"
        path.write_bytes(build_scanned_pdf())
        with pytest.raises(ScannedDocumentError) as info:
            extract(path, name="scan.pdf")
        assert info.value.code == "SCANNED_DOCUMENT_UNSUPPORTED"
        assert info.value.status_code == 422
        assert info.value.details["scanned"] is True
        assert info.value.details["image_only_pages"] == [1]
        assert "OCR" in info.value.message

    def test_an_encrypted_pdf_is_refused(self, tmp_path: Path) -> None:
        path = tmp_path / "locked.pdf"
        path.write_bytes(build_encrypted_pdf())
        with pytest.raises(EncryptedDocumentError) as info:
            extract(path, name="locked.pdf")
        assert info.value.code == "ENCRYPTED_DOCUMENT_UNSUPPORTED"
        assert "password" in info.value.message.lower()

    @pytest.mark.parametrize("payload", [b"%PDF-1.4 not really a pdf", b"just some bytes"])
    def test_a_file_that_is_not_a_pdf_fails_as_an_extraction_error(
        self, tmp_path: Path, payload: bytes
    ) -> None:
        path = tmp_path / "broken.pdf"
        path.write_bytes(payload)
        with pytest.raises(ExtractionError) as info:
            extract(path, name="broken.pdf")
        assert info.value.code == "EXTRACTION_FAILED"
        assert not isinstance(info.value, ScannedDocumentError)

    def test_an_unsupported_kind_is_refused(self, pdf_path: Path) -> None:
        with pytest.raises(UnsupportedFileTypeError) as info:
            extract_document(
                pdf_path,
                file_kind="txt",
                document_id="doc_test",
                document_name="notes.txt",
                max_chars=MAX_CHARS,
            )
        assert info.value.status_code == 415


class TestCharacterBudget:
    def test_a_pdf_is_cut_at_the_budget_and_says_so(self, tmp_path: Path) -> None:
        path = tmp_path / "four.pdf"
        path.write_bytes(build_pdf(pages=4))
        first_page_chars = len(extract(path).segments[0].text)

        document = extract(path, max_chars=first_page_chars + 10)
        assert document.truncated is True
        assert len(document.segments) == 1
        assert any("character budget" in w for w in document.warnings)
        assert any("Increase MAX_EXTRACTED_CHARS" in w for w in document.warnings)

    def test_a_docx_is_cut_at_the_budget_and_says_so(self, tmp_path: Path) -> None:
        path = tmp_path / "report.docx"
        path.write_bytes(build_docx())
        document = extract(path, max_chars=60)
        assert document.truncated is True
        assert len(document.segments) < len(REPORT_DOCX_PARAGRAPHS)
        assert any("character budget" in w for w in document.warnings)

    def test_a_budget_smaller_than_the_first_page_is_reported_not_called_a_scan(
        self, tmp_path: Path
    ) -> None:
        # A text PDF must never be described as a scan just because the budget
        # was too small: that would send the user off to run OCR for nothing.
        path = tmp_path / "four.pdf"
        path.write_bytes(build_pdf(pages=4))
        with pytest.raises(ExtractionBudgetError) as info:
            extract(path, name="four.pdf", max_chars=10)
        assert info.value.code == "EXTRACTION_BUDGET_TOO_SMALL"
        assert info.value.details["max_extracted_chars"] == 10
        assert info.value.details["first_block_chars"] > 10
        assert "MAX_EXTRACTED_CHARS" in info.value.message

    def test_a_budget_smaller_than_the_first_paragraph_is_reported(self, tmp_path: Path) -> None:
        path = tmp_path / "report.docx"
        path.write_bytes(build_docx())
        with pytest.raises(ExtractionBudgetError) as info:
            extract(path, name="report.docx", max_chars=5)
        assert info.value.code == "EXTRACTION_BUDGET_TOO_SMALL"

    def test_zero_disables_the_budget(self, tmp_path: Path) -> None:
        path = tmp_path / "four.pdf"
        path.write_bytes(build_pdf(pages=4))
        document = extract(path, max_chars=0)
        assert document.truncated is False
        assert len(document.segments) == 4


class TestDocxExtraction:
    def test_metadata_describes_the_file(self, docx_path: Path) -> None:
        document = extract(docx_path, name="report.docx")
        assert document.file_kind == "docx"
        assert document.page_count is None
        assert document.has_text_layer is True
        assert document.char_count == sum(len(s.text) for s in document.segments)
        assert document.size_bytes == docx_path.stat().st_size

    def test_body_paragraphs_are_addressable(self, docx_path: Path) -> None:
        document = extract(docx_path)
        paths = [s.paragraph_path for s in document.segments]
        assert paths[0] == "body:p0"
        for index, expected in enumerate(REPORT_DOCX_PARAGRAPHS):
            assert document.segments[index].text == expected

    def test_table_cells_get_row_and_column_paths(self, docx_path: Path) -> None:
        document = extract(docx_path)
        by_path = {s.paragraph_path: s.text for s in document.segments}
        assert by_path["body:tbl0/r0c0:p0"] == TABLE_ROW[0]
        assert by_path["body:tbl0/r0c1:p0"] == TABLE_ROW[1]

    def test_section_headers_are_extracted(self, docx_path: Path) -> None:
        document = extract(docx_path)
        by_path = {s.paragraph_path: s.text for s in document.segments}
        assert by_path["sec0.header:p0"] == "Confidential: Alice Johnson"

    def test_run_spans_reconstruct_the_paragraph_text(self, docx_path: Path) -> None:
        for segment in extract(docx_path).segments:
            assert segment.runs, f"{segment.paragraph_path} has no run spans"
            rebuilt = "".join(segment.text[r.char_start : r.char_end] for r in segment.runs)
            assert rebuilt == segment.text

    def test_run_spans_do_not_overlap(self, docx_path: Path) -> None:
        for segment in extract(docx_path).segments:
            assert segment.runs[0].char_start == 0
            assert segment.runs[-1].char_end == len(segment.text)
            for previous, current in zip(segment.runs, segment.runs[1:]):
                assert current.char_start == previous.char_end

    def test_docx_segments_have_no_page_number(self, docx_path: Path) -> None:
        assert all(s.page is None for s in extract(docx_path).segments)

    def test_limitations_are_declared(self, docx_path: Path) -> None:
        document = extract(docx_path)
        assert document.limitations == DOCX_LIMITATIONS
        assert any("text boxes" in line for line in document.limitations)

    def test_a_value_split_across_runs_is_extracted_as_one_string(self, tmp_path: Path) -> None:
        path = tmp_path / "split.docx"
        path.write_bytes(build_docx_with_split_runs())
        document = extract(path)
        assert document.segments[0].text == "Contact Alice Johnson today."
        # Six runs were flattened into one segment; the run spans still address
        # each of them, which is what makes run-level redaction possible.
        assert len(document.segments[0].runs) == 6

    def test_hyperlink_runs_are_not_missed(self, tmp_path: Path) -> None:
        # paragraph.runs omits hyperlink content, which would leave redactable
        # text behind.
        document = docx.Document()
        paragraph = document.add_paragraph()
        relation = paragraph.part.relate_to(
            "https://example-school.test/alice", RT.HYPERLINK, is_external=True
        )
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), relation)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "a.johnson@example-school.test"
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        document.add_paragraph("Plain paragraph.")
        path = tmp_path / "link.docx"
        document.save(str(path))

        extracted = extract(path)
        assert [s.text for s in extracted.segments] == [
            "a.johnson@example-school.test",
            "Plain paragraph.",
        ]

    def test_a_merged_cell_is_walked_once(self, tmp_path: Path) -> None:
        document = docx.Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "Guardian: Alice Johnson"
        table.cell(1, 0).text = "Phone: 555-222-3333"
        table.cell(1, 1).text = "Course: Algebra II"
        path = tmp_path / "merged.docx"
        document.save(str(path))

        extracted = extract(path)
        paths = [s.paragraph_path for s in extracted.segments]
        assert paths == ["body:tbl0/r0c0:p0", "body:tbl0/r1c0:p0", "body:tbl0/r1c1:p0"]
        assert len(paths) == len(set(paths))

    def test_a_nested_table_is_walked(self, tmp_path: Path) -> None:
        document = docx.Document()
        outer = document.add_table(rows=1, cols=1)
        inner = outer.cell(0, 0).add_table(rows=1, cols=1)
        inner.cell(0, 0).text = "Nested: Alice Johnson"
        outer.cell(0, 0).paragraphs[0].text = "Outer text"
        path = tmp_path / "nested.docx"
        document.save(str(path))

        extracted = extract(path)
        assert [s.text for s in extracted.segments] == [
            "Outer text",
            "Nested: Alice Johnson",
        ]
        assert extracted.segments[1].paragraph_path == "body:tbl0/r0c0:tbl0/r0c0:p0"

    def test_a_document_with_no_readable_text_is_refused(self, tmp_path: Path) -> None:
        path = tmp_path / "blank.docx"
        docx.Document().save(str(path))
        with pytest.raises(EmptyDocumentError) as info:
            extract(path, name="blank.docx")
        assert info.value.code == "EMPTY_DOCUMENT"
        assert "text boxes" in info.value.message

    def test_a_file_that_is_not_a_docx_package_is_refused(self, tmp_path: Path) -> None:
        path = tmp_path / "broken.docx"
        path.write_bytes(b"not a zip file at all")
        with pytest.raises(ExtractionError) as info:
            extract(path, name="broken.docx")
        assert info.value.code == "EXTRACTION_FAILED"
        assert "not a valid .docx package" in info.value.message

    def test_a_zip_that_is_not_a_word_document_is_refused(self, tmp_path: Path) -> None:
        import zipfile

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("hello.txt", "not a word document")
        path = tmp_path / "fake.docx"
        path.write_bytes(buffer.getvalue())
        with pytest.raises(ExtractionError) as info:
            extract(path, name="fake.docx")
        assert info.value.code == "EXTRACTION_FAILED"


class TestUncoveredWordConstructs:
    def test_a_document_with_no_uncovered_constructs_has_no_warning(self, docx_path: Path) -> None:
        assert extract(docx_path).warnings == []

    def test_a_text_box_is_reported_as_uncovered(self, tmp_path: Path) -> None:
        document = docx.Document()
        document.add_paragraph("Visible text: Alice Johnson")
        paragraph = document.add_paragraph()
        # A VML text box: its text is not in a run this release can rewrite, so it
        # must be reported rather than silently skipped.
        paragraph._p.append(
            parse_xml(
                "<w:pict %s>"
                '<v:shape xmlns:v="urn:schemas-microsoft-com:vml" id="tb1">'
                "<v:textbox><w:txbxContent>"
                "<w:p><w:r><w:t>Hidden: 555-111-2222</w:t></w:r></w:p>"
                "</w:txbxContent></v:textbox>"
                "</v:shape></w:pict>" % nsdecls("w")
            )
        )
        path = tmp_path / "textbox.docx"
        document.save(str(path))

        extracted = extract(path)
        assert any("floating text boxes" in w for w in extracted.warnings)
        assert any("not redacted by this release" in w for w in extracted.warnings)
        assert "Hidden: 555-111-2222" not in "\n".join(s.text for s in extracted.segments)
        assert any(s.text == "Visible text: Alice Johnson" for s in extracted.segments)

    def test_a_tracked_change_insertion_is_reported(self, tmp_path: Path) -> None:
        document = docx.Document()
        paragraph = document.add_paragraph()
        insertion = OxmlElement("w:ins")
        # Word always writes these attributes; the detector must not depend on them.
        insertion.set(qn("w:id"), "1")
        insertion.set(qn("w:author"), "Test Editor")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Alice Johnson"
        run.append(text)
        insertion.append(run)
        paragraph._p.append(insertion)
        document.add_paragraph("Untouched paragraph.")
        path = tmp_path / "tracked.docx"
        document.save(str(path))

        extracted = extract(path)
        assert any("tracked-change insertions" in w for w in extracted.warnings)
        assert any("not redacted by this release" in w for w in extracted.warnings)
        # The inserted text is not in a normal run, so it is not in a segment
        # either - which is exactly why the warning exists.
        assert "Alice Johnson" not in "\n".join(s.text for s in extracted.segments)

    def test_a_tracked_change_without_attributes_is_still_reported(self, tmp_path: Path) -> None:
        document = docx.Document()
        paragraph = document.add_paragraph()
        insertion = OxmlElement("w:ins")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Bare insertion"
        run.append(text)
        insertion.append(run)
        paragraph._p.append(insertion)
        document.add_paragraph("Untouched paragraph.")
        path = tmp_path / "bare_ins.docx"
        document.save(str(path))

        assert any(
            "tracked-change insertions" in w for w in extract(path).warnings
        )
