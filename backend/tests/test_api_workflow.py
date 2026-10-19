"""End-to-end workflow through the HTTP API, in deterministic mode.

These tests drive the whole product path - upload, analyse, review, redact,
download - and assert the properties that matter: originals untouched, approved
text actually gone, unapproved text still there.
"""

from __future__ import annotations

import pymupdf
import pytest

from .conftest import (
    INSTRUCTION,
    TABLE_ROW,
    build_docx,
    build_pdf,
    upload,
)


def _ids_matching(body: dict, needle: str) -> list[str]:
    return [
        detection["detection_id"]
        for detection in body["detections"]
        if needle.lower() in detection["matched_text"].lower()
    ]


def _all_ids(body: dict) -> list[str]:
    return [detection["detection_id"] for detection in body["detections"]]


def _spans(body: dict) -> list[tuple]:
    return [
        (
            detection["document_name"],
            detection["page"],
            detection["paragraph_path"],
            detection["char_start"],
            detection["char_end"],
            detection["matched_text"],
        )
        for detection in body["detections"]
    ]


# --------------------------------------------------------------------- analyze
def test_analyze_returns_a_reviewable_plan_without_touching_the_source(
    client, pdf_bytes, tmp_path
):
    original = tmp_path / "original.pdf"
    original.write_bytes(pdf_bytes)
    before = original.read_bytes()

    response = client.post(
        "/api/v1/analyze",
        files=[upload("report.pdf", pdf_bytes)],
        data={"instruction": INSTRUCTION},
    )
    assert response.status_code == 200, response.text
    body = response.json()

    assert body["status"] == "analysed"
    assert body["session_id"].startswith("sess_")
    assert body["summary"]["documents_analysed"] == 1
    assert body["summary"]["detections"] > 0
    assert body["interpretation"]["model_used"] is False
    assert body["interpretation"]["interpretation_mode"] == "deterministic"

    matched = {detection["matched_text"] for detection in body["detections"]}
    assert "Alice Johnson" in matched
    assert "(555) 123-4567" in matched

    # The keep-clause protected the course/grade lines: nothing from them appears.
    assert not any("Algebra" in text for text in matched)

    for detection in body["detections"]:
        assert detection["detection_id"].startswith("det_")
        assert detection["page"] == 1
        assert detection["document_name"] == "report.pdf"

    # Re-analysing the same bytes finds the same spans. Ids differ because each
    # upload gets its own opaque document id, so an id from one analysis is never
    # accepted by another (see test_review_rejects_an_id_from_another_session).
    second = client.post(
        "/api/v1/analyze",
        files=[upload("report.pdf", pdf_bytes)],
        data={"instruction": INSTRUCTION},
    ).json()
    assert _spans(second) == _spans(body)

    assert original.read_bytes() == before


def test_detection_ids_are_a_pure_function_of_the_span():
    from app.matching.engine import detection_id

    kwargs = {
        "document_id": "doc_fixed",
        "segment_index": 3,
        "char_start": 10,
        "char_end": 24,
        "matched_text": "Alice Johnson",
        "strategy": "pattern",
        "pattern_name": "person_name",
    }
    assert detection_id(**kwargs) == detection_id(**kwargs)
    assert detection_id(**kwargs).startswith("det_")
    assert len(detection_id(**kwargs)) == 28
    moved = dict(kwargs, char_start=11)
    assert detection_id(**moved) != detection_id(**kwargs)


def test_analyze_reports_what_it_cannot_do(client, analyzed):
    limitations = " ".join(analyzed["limitations"]).lower()
    assert "not a legal-compliance tool" in limitations
    assert analyzed["interpretation"]["exclusions"]
    assert any(
        "grade" in exclusion["term"].lower() or "course" in exclusion["term"].lower()
        for exclusion in analyzed["interpretation"]["exclusions"]
    )


def test_analyze_handles_multiple_files(client, pdf_bytes, docx_bytes):
    response = client.post(
        "/api/v1/analyze",
        files=[
            upload("report.pdf", pdf_bytes),
            upload("report.docx", docx_bytes),
        ],
        data={"instruction": INSTRUCTION},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["summary"]["documents_uploaded"] == 2
    assert body["summary"]["documents_analysed"] == 2
    kinds = {document["file_kind"] for document in body["documents"]}
    assert kinds == {"pdf", "docx"}
    assert {d["file_kind"] for d in body["detections"]} == {"pdf", "docx"}


def test_docx_table_cells_are_detected(client):
    response = client.post(
        "/api/v1/analyze",
        files=[upload("report.docx", build_docx(table_row=TABLE_ROW))],
        data={"instruction": INSTRUCTION},
    )
    assert response.status_code == 200, response.text
    paths = {d["paragraph_path"] for d in response.json()["detections"]}
    assert "body:tbl0/r0c0:p0" in paths
    assert "body:tbl0/r0c1:p0" in paths


# ----------------------------------------------------------- analyze failures
def test_analyze_rejects_an_empty_instruction(client, pdf_bytes):
    response = client.post(
        "/api/v1/analyze",
        files=[upload("report.pdf", pdf_bytes)],
        data={"instruction": "   "},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "EMPTY_INSTRUCTION"


def test_analyze_rejects_unsupported_extension(client, pdf_bytes):
    response = client.post(
        "/api/v1/analyze",
        files=[upload("report.txt", pdf_bytes)],
        data={"instruction": INSTRUCTION},
    )
    assert response.status_code == 415
    assert response.json()["error"]["code"] == "UNSUPPORTED_FILE_TYPE"


def test_analyze_rejects_a_docx_renamed_to_pdf(client, docx_bytes):
    response = client.post(
        "/api/v1/analyze",
        files=[upload("sneaky.pdf", docx_bytes)],
        data={"instruction": INSTRUCTION},
    )
    assert response.status_code == 415
    error = response.json()["error"]
    assert error["code"] == "FILE_CONTENT_MISMATCH"
    assert error["details"]["detected"] == "zip"


def test_analyze_rejects_a_random_binary_named_docx(client):
    response = client.post(
        "/api/v1/analyze",
        files=[upload("junk.docx", b"\x00\x01not a document at all")],
        data={"instruction": INSTRUCTION},
    )
    assert response.status_code == 415
    assert response.json()["error"]["code"] == "FILE_CONTENT_MISMATCH"


def test_analyze_rejects_too_many_files(client, pdf_bytes):
    response = client.post(
        "/api/v1/analyze",
        files=[upload(f"report{i}.pdf", pdf_bytes) for i in range(4)],
        data={"instruction": INSTRUCTION},
    )
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "TOO_MANY_FILES"


def test_analyze_rejects_an_oversized_file(client, pdf_bytes, tmp_path):
    # settings.max_file_size_bytes is 2 MB in the test profile
    padded = pdf_bytes + b"%PAD\n" + b"x" * (2 * 1024 * 1024 + 1024)
    response = client.post(
        "/api/v1/analyze",
        files=[upload("big.pdf", padded)],
        data={"instruction": INSTRUCTION},
    )
    assert response.status_code == 413
    assert response.json()["error"]["code"] == "FILE_TOO_LARGE"


def test_scanned_pdf_is_reported_not_silently_skipped(client):
    from .conftest import build_scanned_pdf

    response = client.post(
        "/api/v1/analyze",
        files=[upload("scan.pdf", build_scanned_pdf())],
        data={"instruction": INSTRUCTION},
    )
    assert response.status_code == 422
    error = response.json()["error"]
    assert error["code"] == "SCANNED_DOCUMENT_UNSUPPORTED"
    assert error["details"]["scanned"] is True
    assert "OCR" in error["message"]


def test_scanned_pdf_alongside_a_good_file_fails_only_the_scan(client, pdf_bytes):
    from .conftest import build_scanned_pdf

    response = client.post(
        "/api/v1/analyze",
        files=[
            upload("scan.pdf", build_scanned_pdf()),
            upload("report.pdf", pdf_bytes),
        ],
        data={"instruction": INSTRUCTION},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["summary"]["documents_analysed"] == 1
    assert body["summary"]["documents_failed"] == 1
    failure = body["document_failures"][0]
    assert failure["document_name"] == "scan.pdf"
    assert failure["code"] == "SCANNED_DOCUMENT_UNSUPPORTED"
    assert any("could NOT be analysed" in warning for warning in body["warnings"])


def test_encrypted_pdf_is_refused(client):
    from .conftest import build_encrypted_pdf

    response = client.post(
        "/api/v1/analyze",
        files=[upload("locked.pdf", build_encrypted_pdf())],
        data={"instruction": INSTRUCTION},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "ENCRYPTED_DOCUMENT_UNSUPPORTED"


def test_unusable_instruction_returns_examples(client, pdf_bytes):
    response = client.post(
        "/api/v1/analyze",
        files=[upload("report.pdf", pdf_bytes)],
        data={"instruction": "please make it nice and clean"},
    )
    assert response.status_code == 422
    error = response.json()["error"]
    assert error["code"] == "INSTRUCTION_AMBIGUOUS"
    assert error["details"]["supported_detectors"]


def test_analyze_without_files_is_rejected(client):
    response = client.post("/api/v1/analyze", data={"instruction": INSTRUCTION})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "VALIDATION_ERROR"


# ---------------------------------------------------------------------- review
def test_review_records_the_selection(client, analyzed):
    session_id = analyzed["session_id"]
    keep = _ids_matching(analyzed, "Alice Johnson")
    drop = [i for i in _all_ids(analyzed) if i not in keep]

    response = client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": keep, "rejected_detection_ids": drop},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "reviewed"
    assert body["approved"] == len(keep)
    assert body["rejected"] == len(drop)
    assert body["ignored"] == 0
    assert body["per_document"][0]["document_name"] == "report.pdf"


def test_review_warns_about_undecided_detections(client, analyzed):
    session_id = analyzed["session_id"]
    keep = _ids_matching(analyzed, "Alice Johnson")[:1]

    response = client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": keep},
    )
    assert response.status_code == 200
    assert any("will NOT be redacted" in w for w in response.json()["warnings"])


def test_review_rejects_a_malformed_id(client, analyzed):
    response = client.post(
        f"/api/v1/sessions/{analyzed['session_id']}/review",
        json={"approved_detection_ids": ["not-a-detection-id"]},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "INVALID_DETECTION_SELECTION"


def test_review_rejects_a_stale_id(client, analyzed):
    stale = "det_" + "a" * 24
    response = client.post(
        f"/api/v1/sessions/{analyzed['session_id']}/review",
        json={"approved_detection_ids": [stale]},
    )
    assert response.status_code == 409
    error = response.json()["error"]
    assert error["code"] == "STALE_DETECTION_ID"
    assert error["details"]["unknown_ids"] == [stale]


def test_review_rejects_an_id_from_another_session(client, analyzed, pdf_bytes):
    other = client.post(
        "/api/v1/analyze",
        files=[upload("other.pdf", build_pdf(["Different: Bob Smith"]))],
        data={"instruction": "Remove all person names"},
    ).json()
    foreign = _all_ids(other)[0]

    response = client.post(
        f"/api/v1/sessions/{analyzed['session_id']}/review",
        json={"approved_detection_ids": [foreign]},
    )
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "STALE_DETECTION_ID"


def test_review_rejects_a_conflicting_selection(client, analyzed):
    one = _all_ids(analyzed)[0]
    response = client.post(
        f"/api/v1/sessions/{analyzed['session_id']}/review",
        json={"approved_detection_ids": [one], "rejected_detection_ids": [one]},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "INVALID_DETECTION_SELECTION"


def test_unknown_session_is_expired_not_found(client):
    response = client.post(
        "/api/v1/sessions/sess_doesnotexist/review",
        json={"approved_detection_ids": []},
    )
    assert response.status_code == 410
    assert response.json()["error"]["code"] == "SESSION_EXPIRED"

    response = client.post(
        "/api/v1/sessions/garbage/review", json={"approved_detection_ids": []}
    )
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "SESSION_NOT_FOUND"


# ---------------------------------------------------------------------- redact
def test_redact_requires_a_review_first(client, analyzed):
    response = client.post(f"/api/v1/sessions/{analyzed['session_id']}/redact")
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "REVIEW_REQUIRED"


def test_redact_refuses_when_nothing_was_approved(client, analyzed):
    session_id = analyzed["session_id"]
    client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": [], "rejected_detection_ids": _all_ids(analyzed)},
    )
    response = client.post(f"/api/v1/sessions/{session_id}/redact")
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "NO_APPROVED_REDACTIONS"


def test_pdf_redaction_removes_the_text_layer(client, analyzed, tmp_path):
    session_id = analyzed["session_id"]
    approved = _ids_matching(analyzed, "Alice Johnson")
    retained = [i for i in _all_ids(analyzed) if i not in approved]

    client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": approved, "rejected_detection_ids": retained},
    )
    response = client.post(f"/api/v1/sessions/{session_id}/redact")
    assert response.status_code == 200, response.text
    body = response.json()

    assert body["status"] == "redacted"
    assert body["residual_text_found"] is False
    assert len(body["outputs"]) == 1
    assert body["outputs"][0]["output_name"] == "report.redacted.pdf"

    download = client.get(body["outputs"][0]["download_url"])
    assert download.status_code == 200
    saved = tmp_path / "downloaded.pdf"
    saved.write_bytes(download.content)

    with pymupdf.open(saved) as redacted:
        text = "\n".join(page.get_text() for page in redacted)
        assert "Alice Johnson" not in text
        # The unapproved phone number must survive - redaction is opt-in per match.
        assert "(555) 123-4567" in text
        # Unrelated content is preserved.
        assert "Algebra II" in text
        assert redacted.metadata["author"] in ("", None)

    verification = body["verification"][0]
    assert verification["coverage"] == "text-layer-only"
    assert {r["status"] for r in verification["results"]} == {"removed"}


def test_pdf_redaction_is_not_a_black_rectangle(client, analyzed, tmp_path):
    """The glyphs must be gone, not merely covered."""
    session_id = analyzed["session_id"]
    approved = _all_ids(analyzed)
    client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": approved},
    )
    body = client.post(f"/api/v1/sessions/{session_id}/redact").json()
    saved = tmp_path / "all.pdf"
    saved.write_bytes(client.get(body["outputs"][0]["download_url"]).content)

    with pymupdf.open(saved) as redacted:
        for detection in analyzed["detections"]:
            needle = detection["matched_text"]
            assert not any(page.search_for(needle) for page in redacted), needle
        text = " ".join(page.get_text() for page in redacted)

    # The instruction never mentioned email addresses, so that one is untouched:
    # redaction only ever removes what was asked for and approved.
    assert "a.johnson@example-school.test" in text
    assert "Algebra II" in text


def test_docx_redaction_replaces_text_including_tables(client, tmp_path):
    analyzed = client.post(
        "/api/v1/analyze",
        files=[upload("report.docx", build_docx(table_row=TABLE_ROW))],
        data={"instruction": INSTRUCTION},
    ).json()
    session_id = analyzed["session_id"]

    guardian = [
        d["detection_id"]
        for d in analyzed["detections"]
        if d["paragraph_path"] == "body:tbl0/r0c0:p0"
    ]
    assert guardian
    retained = [i for i in _all_ids(analyzed) if i not in guardian]

    client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": guardian, "rejected_detection_ids": retained},
    )
    body = client.post(f"/api/v1/sessions/{session_id}/redact").json()
    assert body["outputs"][0]["output_name"] == "report.redacted.docx"

    saved = tmp_path / "downloaded.docx"
    saved.write_bytes(client.get(body["outputs"][0]["download_url"]).content)

    import docx

    document = docx.Document(str(saved))
    table_text = document.tables[0].rows[0].cells[0].text
    assert "Alice Johnson" not in table_text
    assert "[REDACTED]" in table_text
    # The rest of the paragraph in the neighbouring cell is untouched.
    assert document.tables[0].rows[0].cells[1].text == "Phone: 555-222-3333"
    assert document.core_properties.author == ""


def test_docx_redaction_handles_a_value_split_across_runs(client, tmp_path):
    from .conftest import build_docx_with_split_runs

    analyzed = client.post(
        "/api/v1/analyze",
        files=[upload("split.docx", build_docx_with_split_runs())],
        data={"instruction": "Remove all person names but keep everything else."},
    ).json()
    assert analyzed["summary"]["detections"] >= 1
    session_id = analyzed["session_id"]
    client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": _all_ids(analyzed)},
    )
    body = client.post(f"/api/v1/sessions/{session_id}/redact").json()
    saved = tmp_path / "split.docx"
    saved.write_bytes(client.get(body["outputs"][0]["download_url"]).content)

    import docx

    document = docx.Document(str(saved))
    assert "Alice Johnson" not in document.paragraphs[0].text
    assert document.paragraphs[1].text == "Nothing sensitive here."


def test_redaction_only_applies_approved_matches(client, analyzed, tmp_path):
    session_id = analyzed["session_id"]
    approved = _ids_matching(analyzed, "555-987-6543")
    assert approved

    client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": approved},
    )
    body = client.post(f"/api/v1/sessions/{session_id}/redact").json()
    saved = tmp_path / "one.pdf"
    saved.write_bytes(client.get(body["outputs"][0]["download_url"]).content)

    with pymupdf.open(saved) as redacted:
        text = "\n".join(page.get_text() for page in redacted)
        assert "555-987-6543" not in text
        assert "Alice Johnson" in text

    verification = body["verification"][0]
    assert verification["residual_text_found"] is False


def test_verification_distinguishes_retained_by_choice(client, analyzed, tmp_path):
    """Two identical strings, one approved: the survivor is not a failure."""
    duplicate = build_pdf(
        ["Call Alice Johnson on 555-111-2222.", "Also call Alice Johnson on 555-111-2222."]
    )
    body = client.post(
        "/api/v1/analyze",
        files=[upload("dup.pdf", duplicate)],
        data={"instruction": "Remove every person name."},
    ).json()
    names = [
        detection
        for detection in body["detections"]
        if detection["matched_text"] == "Alice Johnson"
    ]
    assert len(names) == 2, [d["matched_text"] for d in body["detections"]]

    session_id = body["session_id"]
    client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": [names[0]["detection_id"]]},
    )
    redacted = client.post(f"/api/v1/sessions/{session_id}/redact").json()
    report = redacted["verification"][0]
    assert redacted["residual_text_found"] is False

    entry = next(r for r in report["results"] if r["text"] == "Alice Johnson")
    assert entry["status"] == "partially_removed"
    assert entry["approved_count"] == 1
    assert entry["retained_by_choice"] == 1
    assert entry["occurrences_in_output"] == 1


def test_redact_twice_is_refused(client, analyzed):
    session_id = analyzed["session_id"]
    client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": _all_ids(analyzed)},
    )
    first = client.post(f"/api/v1/sessions/{session_id}/redact")
    assert first.status_code == 200
    second = client.post(f"/api/v1/sessions/{session_id}/redact")
    assert second.status_code == 409
    assert second.json()["error"]["code"] == "ALREADY_REDACTED"


def test_originals_are_never_modified(client, analyzed, tmp_path):
    source = tmp_path / "source.pdf"
    source.write_bytes(build_pdf())
    before = source.read_bytes()

    client.post(
        "/api/v1/analyze",
        files=[upload("source.pdf", source.read_bytes())],
        data={"instruction": INSTRUCTION},
    )
    assert source.read_bytes() == before


# -------------------------------------------------------------------- downloads
def test_download_rejects_an_unknown_token(client):
    response = client.get("/api/v1/downloads/dl_nonsense")
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "DOWNLOAD_NOT_FOUND"


def test_download_token_is_opaque_and_path_escape_is_blocked(client, analyzed):
    session_id = analyzed["session_id"]
    client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": _all_ids(analyzed)},
    )
    body = client.post(f"/api/v1/sessions/{session_id}/redact").json()
    token = body["outputs"][0]["download_token"]
    assert token.startswith("dl_")
    assert len(token) > 20

    for attempt in ("..%2f..%2fetc%2fpasswd", token + "%00.pdf", ".."):
        response = client.get(f"/api/v1/downloads/{attempt}")
        assert response.status_code in (404, 422), attempt


def test_expired_download_is_gone(client, analyzed, settings):
    session_id = analyzed["session_id"]
    client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": _all_ids(analyzed)},
    )
    body = client.post(f"/api/v1/sessions/{session_id}/redact").json()
    output = body["outputs"][0]

    from datetime import timedelta

    from app.storage import _now

    store = client.app.state.store
    live = store.get(session_id)
    live.outputs[0].expires_at = _now() - timedelta(minutes=1)

    response = client.get(output["download_url"])
    assert response.status_code == 410
    assert response.json()["error"]["code"] == "DOWNLOAD_EXPIRED"


# --------------------------------------------------------------------- sessions
def test_session_status_and_delete(client, analyzed):
    session_id = analyzed["session_id"]
    status = client.get(f"/api/v1/sessions/{session_id}")
    assert status.status_code == 200
    assert status.json()["status"] == "analysed"
    assert status.json()["detections"] == analyzed["summary"]["detections"]

    deleted = client.delete(f"/api/v1/sessions/{session_id}")
    assert deleted.status_code == 200
    assert deleted.json()["deleted"] is True

    assert client.get(f"/api/v1/sessions/{session_id}").status_code == 410
    assert client.delete(f"/api/v1/sessions/{session_id}").json()["deleted"] is False


@pytest.mark.parametrize(
    "instruction, expected_present, expected_absent",
    [
        ("Remove all email addresses.", ["Alice Johnson"], ["a.johnson@example-school.test"]),
        (
            "Redact the phrase 'Springfield High' but keep everything else.",
            ["Alice Johnson", "(555) 123-4567"],
            ["Springfield High"],
        ),
    ],
)
def test_other_instruction_shapes(
    client, pdf_bytes, tmp_path, instruction, expected_present, expected_absent
):
    analyzed = client.post(
        "/api/v1/analyze",
        files=[upload("report.pdf", pdf_bytes)],
        data={"instruction": instruction},
    ).json()
    assert analyzed["summary"]["detections"] >= 1, analyzed["warnings"]

    session_id = analyzed["session_id"]
    client.post(
        f"/api/v1/sessions/{session_id}/review",
        json={"approved_detection_ids": _all_ids(analyzed)},
    )
    body = client.post(f"/api/v1/sessions/{session_id}/redact").json()
    saved = tmp_path / "out.pdf"
    saved.write_bytes(client.get(body["outputs"][0]["download_url"]).content)

    with pymupdf.open(saved) as redacted:
        text = " ".join(page.get_text() for page in redacted)
    text = " ".join(text.split())
    for needle in expected_absent:
        assert needle not in text, needle
    for needle in expected_present:
        assert needle in text, needle
