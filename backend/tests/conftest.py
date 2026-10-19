"""Shared fixtures.

Every document used here is generated in-memory from obviously synthetic data.
No real personal document is ever read, and nothing is written outside a pytest
``tmp_path``.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Iterator
from pathlib import Path

import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient

from app.config import Settings
from app.extraction.base import ExtractedDocument, Segment
from app.main import create_app
from app.matching.engine import MatchPlan, Target, detect
from app.storage import SessionStore

# ---------------------------------------------------------------------------
# Synthetic content. Names, numbers and addresses here are fabricated.
# ---------------------------------------------------------------------------
REPORT_LINES = [
    "Springfield High Report Card",
    "Student: Alice Johnson",
    "Phone: (555) 123-4567",
    "Home Address: 42 Maple Street, Springfield, IL 62704",
    "Course: Algebra II Grade: A",
    "Course: English Literature Grade: B+",
    "Emergency contact: Brian O'Connor 555-987-6543",
    "Counsellor email: a.johnson@example-school.test",
]

REPORT_DOCX_PARAGRAPHS = REPORT_LINES[:7]

TABLE_ROW = ("Guardian: Alice Johnson", "Phone: 555-222-3333")

INSTRUCTION = (
    "Remove all student names, phone numbers and home addresses, "
    "but keep grades and course names unchanged."
)


def build_pdf(lines: list[str] | None = None, *, pages: int = 1) -> bytes:
    """A minimal text PDF. ``pages > 1`` repeats the block on each page."""
    lines = REPORT_LINES if lines is None else lines
    document = pymupdf.open()
    for _ in range(max(1, pages)):
        page = document.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 560, 780), "\n".join(lines), fontsize=11)
    data = document.tobytes()
    document.close()
    return data


def build_scanned_pdf() -> bytes:
    """A PDF with an image and no text layer at all."""
    document = pymupdf.open()
    page = document.new_page()
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 120), 1)
    pixmap.clear_with(200)
    page.insert_image(pymupdf.Rect(50, 50, 250, 170), pixmap=pixmap)
    data = document.tobytes()
    document.close()
    return data


def build_encrypted_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 560, 780), "Secret: Alice Johnson", fontsize=11)
    data = document.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user"
    )
    document.close()
    return data


def build_docx(
    paragraphs: list[str] | None = None,
    *,
    table_row: tuple[str, str] | None = None,
    header_text: str | None = None,
    author: str = "Alice Johnson",
) -> bytes:
    document = docx.Document()
    for text in (REPORT_DOCX_PARAGRAPHS if paragraphs is None else paragraphs):
        document.add_paragraph(text)

    if table_row is not None:
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = table_row[0]
        table.rows[0].cells[1].text = table_row[1]

    if header_text is not None:
        document.sections[0].header.paragraphs[0].text = header_text

    document.core_properties.author = author
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_docx_with_split_runs() -> bytes:
    """A DOCX where one sensitive value is spread across several runs."""
    document = docx.Document()
    paragraph = document.add_paragraph()
    for piece in ("Contact ", "Alice", " ", "John", "son", " today."):
        paragraph.add_run(piece)
    document.add_paragraph("Nothing sensitive here.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload(filename: str, data: bytes, content_type: str = "application/octet-stream"):
    """Multipart tuple for ``TestClient.post(files=[...])``."""
    return ("files", (filename, io.BytesIO(data), content_type))


# ---------------------------------------------------------------------------
# Detection helpers
# ---------------------------------------------------------------------------
def one_segment_document(text: str, *, name: str = "sample.pdf") -> ExtractedDocument:
    return ExtractedDocument(
        document_id="doc_test",
        document_name=name,
        file_kind="pdf",
        size_bytes=len(text),
        sha256="",
        page_count=1,
        char_count=len(text),
        has_text_layer=True,
        segments=[
            Segment(
                segment_index=0,
                text=text,
                page=1,
                line_ranges=[(start, start + len(line)) for start, line in _lines(text)],
            )
        ],
    )


def _lines(text: str) -> list[tuple[int, str]]:
    cursor = 0
    result = []
    for line in text.split("\n"):
        result.append((cursor, line))
        cursor += len(line) + 1
    return result


def document_with_segments(
    texts: list[str],
    *,
    document_id: str = "doc_a",
    name: str = "sample.pdf",
    file_kind: str = "pdf",
    has_text_layer: bool = True,
) -> ExtractedDocument:
    """An ``ExtractedDocument`` with one segment per entry in ``texts``.

    PDF segments are pages (so ``page`` is set); DOCX segments are paragraphs
    (so ``paragraph_path`` is set instead). Line ranges are computed the same way
    the real extractors compute them, because exclusion clause-bounding depends
    on them.
    """
    is_pdf = file_kind == "pdf"
    segments = []
    for index, text in enumerate(texts):
        ranges: list[tuple[int, int]] = []
        cursor = 0
        for line in text.split("\n"):
            ranges.append((cursor, cursor + len(line)))
            cursor += len(line) + 1
        segments.append(
            Segment(
                segment_index=index,
                text=text,
                page=index + 1 if is_pdf else None,
                paragraph_path=None if is_pdf else f"body:p{index}",
                line_ranges=ranges,
            )
        )
    total = sum(len(text) for text in texts)
    return ExtractedDocument(
        document_id=document_id,
        document_name=name,
        file_kind=file_kind,  # type: ignore[arg-type]
        size_bytes=total,
        sha256="",
        page_count=len(texts),
        char_count=total,
        has_text_layer=has_text_layer,
        segments=segments,
    )


def detect_pattern(pattern_name: str, text: str, *, term: str | None = None) -> list[str]:
    """Run one detector over ``text`` and return the matched strings."""
    plan = MatchPlan(
        targets=[
            Target(
                kind="pattern",
                term=term or pattern_name,
                pattern=pattern_name,
                category=pattern_name,
            )
        ]
    )
    result = detect([one_segment_document(text)], plan)
    return [detection.matched_text for detection in result.detections]


def detect_exact(term: str, text: str) -> list[str]:
    plan = MatchPlan(targets=[Target(kind="exact", term=term, category="named text")])
    result = detect([one_segment_document(text)], plan)
    return [detection.matched_text for detection in result.detections]


# ---------------------------------------------------------------------------
# Settings / app
# ---------------------------------------------------------------------------
def make_settings(tmp_path: Path, **overrides) -> Settings:
    """Settings that ignore any real .env and never need the model server."""
    defaults: dict = {
        "environment": "testing",
        "storage_root": tmp_path / "var",
        "log_json": False,
        "log_level": "WARNING",
        "interpretation_mode": "deterministic",
        "session_ttl_minutes": 30,
        "download_ttl_minutes": 15,
        "cleanup_interval_seconds": 3600,
        "max_upload_files": 3,
        "max_file_size_bytes": 2 * 1024 * 1024,
        "max_total_upload_bytes": 4 * 1024 * 1024,
        "max_extracted_chars": 100_000,
        "llm_base_url": "http://127.0.0.1:8080/v1",
        "llm_model": "gemma-3-1b-it",
    }
    defaults.update(overrides)
    return Settings(_env_file=None, **defaults)


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return make_settings(tmp_path)


@pytest.fixture
def store(settings: Settings) -> SessionStore:
    return SessionStore(settings)


@pytest.fixture
def client(settings: Settings) -> Iterator[TestClient]:
    app = create_app(settings)
    with TestClient(app) as test_client:
        yield test_client
        app.state.store.purge_all()


@pytest.fixture
def tolerant_client(settings: Settings) -> Iterator[TestClient]:
    """A client that returns the 500 envelope instead of re-raising."""
    app = create_app(settings)
    with TestClient(app, raise_server_exceptions=False) as test_client:
        yield test_client
        app.state.store.purge_all()


@pytest.fixture
def pdf_bytes() -> bytes:
    return build_pdf()


@pytest.fixture
def docx_bytes() -> bytes:
    return build_docx(table_row=TABLE_ROW)


@pytest.fixture
def analyzed(client: TestClient, pdf_bytes: bytes) -> dict:
    """Run /analyze in deterministic mode and return the JSON body."""
    response = client.post(
        "/api/v1/analyze",
        files=[upload("report.pdf", pdf_bytes)],
        data={"instruction": INSTRUCTION},
    )
    assert response.status_code == 200, response.text
    return response.json()


@pytest.fixture
def make_pdf_file(tmp_path: Path) -> Callable[..., Path]:
    def _make(name: str = "report.pdf", **kwargs) -> Path:
        path = tmp_path / name
        path.write_bytes(build_pdf(**kwargs))
        return path

    return _make


@pytest.fixture
def make_docx_file(tmp_path: Path) -> Callable[..., Path]:
    def _make(name: str = "report.docx", **kwargs) -> Path:
        path = tmp_path / name
        path.write_bytes(build_docx(**kwargs))
        return path

    return _make
